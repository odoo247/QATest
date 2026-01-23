import base64
import csv
import io
import json
import logging
from datetime import datetime

from odoo import models, fields, api, _
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)


class MigrationProject(models.Model):
    _name = 'migration.project'
    _description = 'Data Migration Project'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'create_date desc'

    name = fields.Char(
        string='Project Name',
        required=True,
        tracking=True,
    )
    description = fields.Text(string='Description')
    
    state = fields.Selection([
        ('draft', 'Draft'),
        ('mapping', 'Mapping'),
        ('validated', 'Validated'),
        ('importing', 'Importing'),
        ('done', 'Completed'),
        ('error', 'Error'),
        ('cancelled', 'Cancelled'),
    ], default='draft', string='Status', tracking=True)
    
    # Source Configuration
    source_type = fields.Selection([
        ('csv', 'CSV File'),
        ('excel', 'Excel File'),
        ('json', 'JSON File'),
    ], string='Source Type', required=True, default='csv')
    
    source_file = fields.Binary(
        string='Source File',
        attachment=True,
    )
    source_filename = fields.Char(string='Source Filename')
    
    csv_delimiter = fields.Char(
        string='CSV Delimiter',
        default=',',
        help='Delimiter used in CSV file',
    )
    csv_encoding = fields.Selection([
        ('utf-8', 'UTF-8'),
        ('latin-1', 'Latin-1'),
        ('cp1252', 'Windows-1252'),
    ], string='File Encoding', default='utf-8')
    
    has_header = fields.Boolean(
        string='Has Header Row',
        default=True,
    )
    
    # Target Configuration
    target_model_id = fields.Many2one(
        'ir.model',
        string='Target Model',
        required=True,
        domain=[('transient', '=', False)],
    )
    target_model_name = fields.Char(
        related='target_model_id.model',
        string='Model Technical Name',
        store=True,
    )
    
    # Mappings
    mapping_ids = fields.One2many(
        'migration.mapping',
        'project_id',
        string='Field Mappings',
    )
    
    # Import Settings
    batch_size = fields.Integer(
        string='Batch Size',
        default=100,
        help='Number of records to import per batch',
    )
    on_error = fields.Selection([
        ('stop', 'Stop Import'),
        ('skip', 'Skip Record'),
        ('log', 'Log and Continue'),
    ], string='On Error', default='skip')
    
    create_if_not_exists = fields.Boolean(
        string='Create Related Records',
        default=False,
        help='Create related records if they do not exist',
    )
    update_existing = fields.Boolean(
        string='Update Existing Records',
        default=False,
        help='Update records if they already exist (based on unique field)',
    )
    unique_field_id = fields.Many2one(
        'ir.model.fields',
        string='Unique Identifier Field',
        domain="[('model_id', '=', target_model_id)]",
        help='Field used to identify existing records for updates',
    )
    
    # Statistics
    total_rows = fields.Integer(
        string='Total Rows',
        readonly=True,
    )
    imported_count = fields.Integer(
        string='Imported Records',
        readonly=True,
    )
    error_count = fields.Integer(
        string='Error Count',
        readonly=True,
    )
    skipped_count = fields.Integer(
        string='Skipped Records',
        readonly=True,
    )
    
    # Logging
    log_ids = fields.One2many(
        'migration.log',
        'project_id',
        string='Migration Logs',
    )
    
    # AI Configuration
    ai_provider = fields.Selection([
        ('anthropic', 'Anthropic Claude'),
        ('openai', 'OpenAI GPT'),
    ], string='AI Provider', default='anthropic')
    
    ai_api_key = fields.Char(
        string='AI API Key',
        help='API key for AI service. Leave empty to use system parameter.',
    )
    
    # Template
    template_id = fields.Many2one(
        'migration.template',
        string='Migration Template',
    )
    
    # Source Data Preview
    source_columns = fields.Text(
        string='Source Columns',
        readonly=True,
    )
    preview_data = fields.Text(
        string='Preview Data',
        readonly=True,
    )
    
    # Progress tracking
    progress = fields.Float(
        string='Progress',
        compute='_compute_progress',
    )
    
    @api.depends('total_rows', 'imported_count', 'error_count', 'skipped_count')
    def _compute_progress(self):
        for record in self:
            if record.total_rows:
                processed = record.imported_count + record.error_count + record.skipped_count
                record.progress = (processed / record.total_rows) * 100
            else:
                record.progress = 0
    
    @api.onchange('source_file')
    def _onchange_source_file(self):
        """Parse source file and extract columns when file is uploaded."""
        if self.source_file:
            try:
                self._parse_source_file()
            except Exception as e:
                raise UserError(_('Error parsing file: %s') % str(e))
    
    def _parse_source_file(self):
        """Parse the source file and extract column information."""
        self.ensure_one()
        if not self.source_file:
            return
        
        file_content = base64.b64decode(self.source_file)
        
        if self.source_type == 'csv':
            self._parse_csv(file_content)
        elif self.source_type == 'excel':
            self._parse_excel(file_content)
        elif self.source_type == 'json':
            self._parse_json(file_content)
    
    def _parse_csv(self, file_content):
        """Parse CSV file content."""
        try:
            content = file_content.decode(self.csv_encoding or 'utf-8')
            reader = csv.reader(io.StringIO(content), delimiter=self.csv_delimiter or ',')
            rows = list(reader)
            
            if not rows:
                raise UserError(_('CSV file is empty'))
            
            if self.has_header:
                columns = rows[0]
                data_rows = rows[1:]
            else:
                columns = [f'Column_{i}' for i in range(len(rows[0]))]
                data_rows = rows
            
            self.source_columns = json.dumps(columns)
            self.total_rows = len(data_rows)
            
            # Store preview (first 5 rows)
            preview = []
            for row in data_rows[:5]:
                preview.append(dict(zip(columns, row)))
            self.preview_data = json.dumps(preview, indent=2)
            
        except UnicodeDecodeError:
            raise UserError(_('Cannot decode file with encoding %s. Try a different encoding.') % self.csv_encoding)
    
    def _parse_excel(self, file_content):
        """Parse Excel file content."""
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
            sheet = workbook.active
            
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                raise UserError(_('Excel file is empty'))
            
            if self.has_header:
                columns = [str(c) if c else f'Column_{i}' for i, c in enumerate(rows[0])]
                data_rows = rows[1:]
            else:
                columns = [f'Column_{i}' for i in range(len(rows[0]))]
                data_rows = rows
            
            self.source_columns = json.dumps(columns)
            self.total_rows = len(data_rows)
            
            # Store preview
            preview = []
            for row in data_rows[:5]:
                preview.append(dict(zip(columns, [str(v) if v else '' for v in row])))
            self.preview_data = json.dumps(preview, indent=2)
            
        except ImportError:
            raise UserError(_('openpyxl library is required for Excel files. Install with: pip install openpyxl'))
    
    def _parse_json(self, file_content):
        """Parse JSON file content."""
        try:
            data = json.loads(file_content.decode('utf-8'))
            
            if isinstance(data, list) and data:
                # Array of objects
                columns = list(data[0].keys()) if isinstance(data[0], dict) else []
                self.total_rows = len(data)
                preview = data[:5]
            elif isinstance(data, dict):
                # Single object or nested structure
                if 'data' in data and isinstance(data['data'], list):
                    columns = list(data['data'][0].keys()) if data['data'] else []
                    self.total_rows = len(data['data'])
                    preview = data['data'][:5]
                else:
                    columns = list(data.keys())
                    self.total_rows = 1
                    preview = [data]
            else:
                raise UserError(_('Unsupported JSON structure'))
            
            self.source_columns = json.dumps(columns)
            self.preview_data = json.dumps(preview, indent=2)
            
        except json.JSONDecodeError as e:
            raise UserError(_('Invalid JSON file: %s') % str(e))
    
    def action_parse_file(self):
        """Manually trigger file parsing."""
        self.ensure_one()
        if not self.source_file:
            raise UserError(_('Please upload a source file first.'))
        self._parse_source_file()
        return True
    
    def action_ai_mapping(self):
        """Open AI mapping wizard."""
        self.ensure_one()
        if not self.source_columns:
            raise UserError(_('Please upload and parse a source file first.'))
        
        return {
            'name': _('AI Field Mapping'),
            'type': 'ir.actions.act_window',
            'res_model': 'migration.ai.mapping.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_project_id': self.id,
            },
        }
    
    def action_validate(self):
        """Validate the migration configuration and data."""
        self.ensure_one()
        self._validate_mappings()
        self._validate_data()
        self.state = 'validated'
        self._log_action('validate', 'Migration validated successfully')
        return True
    
    def _validate_mappings(self):
        """Validate field mappings."""
        if not self.mapping_ids:
            raise UserError(_('No field mappings defined. Please configure mappings first.'))
        
        # Check required fields
        required_fields = self.env['ir.model.fields'].search([
            ('model_id', '=', self.target_model_id.id),
            ('required', '=', True),
            ('name', 'not in', ['id', 'create_uid', 'create_date', 'write_uid', 'write_date']),
        ])
        
        mapped_fields = self.mapping_ids.filtered(lambda m: m.target_field_id).mapped('target_field_id')
        missing_required = required_fields - mapped_fields
        
        if missing_required:
            field_names = ', '.join(missing_required.mapped('field_description'))
            raise UserError(_('Missing mappings for required fields: %s') % field_names)
    
    def _validate_data(self):
        """Validate source data against mappings."""
        errors = []
        warnings = []
        
        # Get sample data for validation
        if not self.preview_data:
            return
        
        preview = json.loads(self.preview_data)
        
        for mapping in self.mapping_ids.filtered(lambda m: m.target_field_id):
            field = mapping.target_field_id
            source_col = mapping.source_column
            
            for i, row in enumerate(preview):
                value = row.get(source_col, '')
                
                # Type validation
                if field.ttype == 'integer' and value:
                    try:
                        int(value)
                    except ValueError:
                        errors.append(f"Row {i+1}: '{value}' is not a valid integer for field '{field.field_description}'")
                
                elif field.ttype == 'float' and value:
                    try:
                        float(value)
                    except ValueError:
                        errors.append(f"Row {i+1}: '{value}' is not a valid number for field '{field.field_description}'")
                
                elif field.ttype in ('date', 'datetime') and value:
                    # Will be handled by transformation rules
                    pass
        
        if errors:
            self._log_action('validation_error', '\n'.join(errors[:20]))  # Log first 20 errors
        
        return errors, warnings
    
    def action_start_import(self):
        """Start the import process."""
        self.ensure_one()
        
        if self.state not in ('validated', 'error'):
            raise UserError(_('Please validate the migration before importing.'))
        
        self.state = 'importing'
        self.imported_count = 0
        self.error_count = 0
        self.skipped_count = 0
        
        self._log_action('import_start', 'Import started')
        
        # Run import in background
        self.with_delay()._run_import() if hasattr(self, 'with_delay') else self._run_import()
        
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Import Started'),
                'message': _('The import process has started. You can monitor progress on the project form.'),
                'type': 'info',
                'sticky': False,
            }
        }
    
    def _run_import(self):
        """Execute the data import."""
        self.ensure_one()
        
        try:
            data = self._get_source_data()
            target_model = self.env[self.target_model_name]
            
            batch = []
            for i, row in enumerate(data):
                try:
                    vals = self._transform_row(row)
                    
                    if self.update_existing and self.unique_field_id:
                        existing = target_model.search([
                            (self.unique_field_id.name, '=', vals.get(self.unique_field_id.name))
                        ], limit=1)
                        
                        if existing:
                            existing.write(vals)
                            self.imported_count += 1
                            self._log_action('update', f'Updated record {existing.id}', row_num=i+1)
                            continue
                    
                    batch.append(vals)
                    
                    if len(batch) >= self.batch_size:
                        self._import_batch(target_model, batch, i)
                        batch = []
                        
                except Exception as e:
                    self._handle_import_error(e, i, row)
                    if self.on_error == 'stop':
                        raise
            
            # Import remaining batch
            if batch:
                self._import_batch(target_model, batch, len(data))
            
            self.state = 'done'
            self._log_action('import_complete', f'Import completed. {self.imported_count} records imported, {self.error_count} errors, {self.skipped_count} skipped.')
            
        except Exception as e:
            self.state = 'error'
            self._log_action('import_error', f'Import failed: {str(e)}')
            raise
    
    def _import_batch(self, model, batch, current_row):
        """Import a batch of records."""
        try:
            records = model.create(batch)
            self.imported_count += len(records)
            self._log_action('batch_import', f'Imported batch of {len(records)} records', row_num=current_row)
        except Exception as e:
            # Try individual imports if batch fails
            for vals in batch:
                try:
                    model.create(vals)
                    self.imported_count += 1
                except Exception as individual_e:
                    self._handle_import_error(individual_e, current_row, vals)
    
    def _handle_import_error(self, error, row_num, row_data):
        """Handle import error based on configuration."""
        self.error_count += 1
        self._log_action('error', str(error), row_num=row_num, row_data=json.dumps(row_data) if isinstance(row_data, dict) else str(row_data))
        
        if self.on_error == 'stop':
            raise UserError(_('Import stopped at row %d: %s') % (row_num, str(error)))
    
    def _get_source_data(self):
        """Get all source data as list of dictionaries."""
        if not self.source_file:
            return []
        
        file_content = base64.b64decode(self.source_file)
        columns = json.loads(self.source_columns)
        
        if self.source_type == 'csv':
            content = file_content.decode(self.csv_encoding or 'utf-8')
            reader = csv.reader(io.StringIO(content), delimiter=self.csv_delimiter or ',')
            rows = list(reader)
            if self.has_header:
                rows = rows[1:]
            return [dict(zip(columns, row)) for row in rows]
        
        elif self.source_type == 'excel':
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if self.has_header:
                rows = rows[1:]
            return [dict(zip(columns, [str(v) if v else '' for v in row])) for row in rows]
        
        elif self.source_type == 'json':
            data = json.loads(file_content.decode('utf-8'))
            if isinstance(data, list):
                return data
            elif 'data' in data:
                return data['data']
            return [data]
        
        return []
    
    def _transform_row(self, row):
        """Transform a source row to target model values."""
        vals = {}
        transformer = self.env['migration.data.transformer']
        
        for mapping in self.mapping_ids.filtered(lambda m: m.target_field_id):
            source_value = row.get(mapping.source_column, '')
            
            # Apply transformations
            if mapping.transformation_ids:
                for transform in mapping.transformation_ids:
                    source_value = transformer.apply_transformation(
                        source_value,
                        transform.transformation_type,
                        transform.transformation_params,
                    )
            
            # Handle default value
            if not source_value and mapping.default_value:
                source_value = mapping.default_value
            
            # Skip if still empty and not required
            if not source_value and not mapping.target_field_id.required:
                continue
            
            # Convert to target field type
            converted_value = self._convert_value(
                source_value,
                mapping.target_field_id,
                mapping,
            )
            
            if converted_value is not None:
                vals[mapping.target_field_id.name] = converted_value
        
        return vals
    
    def _convert_value(self, value, field, mapping):
        """Convert value to target field type."""
        if not value and value != 0:
            return None
        
        try:
            if field.ttype == 'boolean':
                return str(value).lower() in ('true', '1', 'yes', 'y', 't')
            
            elif field.ttype == 'integer':
                return int(float(value))
            
            elif field.ttype == 'float':
                return float(value)
            
            elif field.ttype == 'monetary':
                # Remove currency symbols and convert
                clean_value = ''.join(c for c in str(value) if c.isdigit() or c in '.-')
                return float(clean_value) if clean_value else 0.0
            
            elif field.ttype == 'date':
                return self._parse_date(value, mapping.date_format)
            
            elif field.ttype == 'datetime':
                return self._parse_datetime(value, mapping.date_format)
            
            elif field.ttype == 'many2one':
                return self._resolve_many2one(value, field, mapping)
            
            elif field.ttype == 'many2many':
                return self._resolve_many2many(value, field, mapping)
            
            elif field.ttype == 'selection':
                return self._resolve_selection(value, field, mapping)
            
            else:
                return str(value)
                
        except Exception as e:
            _logger.warning(f"Error converting value '{value}' for field {field.name}: {e}")
            return None
    
    def _parse_date(self, value, date_format=None):
        """Parse date string to date object."""
        if not date_format:
            # Try common formats
            formats = ['%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%Y/%m/%d']
        else:
            formats = [date_format]
        
        for fmt in formats:
            try:
                return datetime.strptime(str(value), fmt).date()
            except ValueError:
                continue
        
        raise ValueError(f"Cannot parse date: {value}")
    
    def _parse_datetime(self, value, date_format=None):
        """Parse datetime string to datetime object."""
        if not date_format:
            formats = [
                '%Y-%m-%d %H:%M:%S',
                '%d/%m/%Y %H:%M:%S',
                '%Y-%m-%dT%H:%M:%S',
                '%Y-%m-%d',
            ]
        else:
            formats = [date_format]
        
        for fmt in formats:
            try:
                return datetime.strptime(str(value), fmt)
            except ValueError:
                continue
        
        raise ValueError(f"Cannot parse datetime: {value}")
    
    def _resolve_many2one(self, value, field, mapping):
        """Resolve many2one field value."""
        if not value:
            return False
        
        relation_model = self.env[field.relation]
        search_field = mapping.relation_search_field or 'name'
        
        # Search for existing record
        record = relation_model.search([(search_field, '=', value)], limit=1)
        
        if record:
            return record.id
        elif self.create_if_not_exists:
            # Create new record
            new_record = relation_model.create({search_field: value})
            return new_record.id
        else:
            return False
    
    def _resolve_many2many(self, value, field, mapping):
        """Resolve many2many field value."""
        if not value:
            return [(5, 0, 0)]  # Clear existing
        
        relation_model = self.env[field.relation]
        search_field = mapping.relation_search_field or 'name'
        delimiter = mapping.multi_value_delimiter or ','
        
        values = [v.strip() for v in str(value).split(delimiter)]
        record_ids = []
        
        for val in values:
            record = relation_model.search([(search_field, '=', val)], limit=1)
            if record:
                record_ids.append(record.id)
            elif self.create_if_not_exists:
                new_record = relation_model.create({search_field: val})
                record_ids.append(new_record.id)
        
        return [(6, 0, record_ids)]
    
    def _resolve_selection(self, value, field, mapping):
        """Resolve selection field value."""
        # Get selection options
        selection = field.selection_ids if hasattr(field, 'selection_ids') else []
        
        # Check value mapping
        if mapping.value_mapping:
            value_map = json.loads(mapping.value_mapping)
            if str(value) in value_map:
                return value_map[str(value)]
        
        # Try direct match
        return str(value)
    
    def _log_action(self, action_type, message, row_num=None, row_data=None):
        """Create a migration log entry."""
        self.env['migration.log'].create({
            'project_id': self.id,
            'action_type': action_type,
            'message': message,
            'row_number': row_num,
            'row_data': row_data,
        })
    
    def action_reset_to_draft(self):
        """Reset project to draft state."""
        self.ensure_one()
        self.state = 'draft'
        self.imported_count = 0
        self.error_count = 0
        self.skipped_count = 0
        return True
    
    def action_view_logs(self):
        """View migration logs."""
        return {
            'name': _('Migration Logs'),
            'type': 'ir.actions.act_window',
            'res_model': 'migration.log',
            'view_mode': 'tree,form',
            'domain': [('project_id', '=', self.id)],
            'context': {'default_project_id': self.id},
        }
    
    def action_view_imported_records(self):
        """View imported records."""
        return {
            'name': _('Imported Records'),
            'type': 'ir.actions.act_window',
            'res_model': self.target_model_name,
            'view_mode': 'tree,form',
            'domain': [('create_uid', '=', self.env.uid)],  # Simplified filter
        }
    
    def action_save_as_template(self):
        """Save current configuration as a template."""
        self.ensure_one()
        
        return {
            'name': _('Save as Template'),
            'type': 'ir.actions.act_window',
            'res_model': 'migration.save.template.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_project_id': self.id,
            },
        }
    
    def action_load_template(self):
        """Load configuration from a template."""
        self.ensure_one()
        
        if not self.template_id:
            raise UserError(_('Please select a template first.'))
        
        self.template_id.apply_to_project(self)
        return True
